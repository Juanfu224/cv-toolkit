from __future__ import annotations

import argparse
import sys
from pathlib import Path

from jinja2 import Environment, FileSystemLoader
from weasyprint import CSS, HTML

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import contact_links, load_yaml
from paths import BASE, PLANTILLAS


from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def apply_perfil(cv: dict) -> dict:
    perfil = load_yaml(BASE / "perfil.yaml")
    cv["nombre"] = perfil["nombre"]
    cv["contacto"] = perfil["contacto"]
    return cv


def bullet_text(item) -> str:
    if isinstance(item, dict):
        return item.get("texto") or ""
    return str(item)


def idiomas_linea(cv: dict) -> str:
    parts = []
    for item in cv.get("idiomas") or []:
        if isinstance(item, dict):
            extra = f" ({item['organismo']})" if item.get("organismo") else ""
            parts.append(f"{item.get('idioma')}: {item.get('nivel')}{extra}")
        else:
            parts.append(str(item))
    return " · ".join(parts)


def certs_linea(cv: dict) -> str:
    parts = []
    for item in cv.get("certificaciones") or []:
        if isinstance(item, dict):
            entidad = f" ({item['entidad']})" if item.get("entidad") else ""
            parts.append(f"{item.get('nombre')}{entidad}")
        else:
            parts.append(str(item))
    return " · ".join(parts)


def to_markdown(cv: dict) -> str:
    c = cv["contacto"]
    lines = [
        f"# {cv['nombre']}",
        "",
        f"## {cv['headline']}",
        "",
        f"{c['telefono']} · {c['email']} · {c['ciudad']}, {c['pais']}",
    ]
    for url in contact_links(c):
        lines.append(url)
    lines.append("")
    for seccion in cv.get("secciones") or []:
        if seccion == "perfil":
            lines += ["## Perfil profesional", "", cv.get("perfil", ""), ""]
        elif seccion == "competencias":
            lines += ["## Competencias", "", " · ".join(cv.get("competencias") or []), ""]
        elif seccion == "experiencia":
            lines += ["## Experiencia", ""]
            for rol in cv.get("experiencia") or []:
                lines.append(f"### {rol['titulo']} — {rol['empresa']} ({rol['fechas']})")
                lines.append("")
                for b in rol.get("bullets") or []:
                    lines.append(f"- {bullet_text(b)}")
                lines.append("")
        elif seccion == "proyectos" and cv.get("proyectos"):
            lines += ["## Proyectos", ""]
            for proj in cv["proyectos"]:
                url = f" — {proj['url']}" if proj.get("url") else ""
                lines.append(f"### {proj['nombre']}{url}")
                lines.append("")
                if proj.get("stack"):
                    lines.append(", ".join(proj["stack"]))
                    lines.append("")
                for b in proj.get("bullets") or []:
                    lines.append(f"- {bullet_text(b)}")
                lines.append("")
        elif seccion == "formacion":
            lines += ["## Formación", ""]
            for edu in cv.get("formacion") or []:
                centro = f" — {edu['centro']}" if edu.get("centro") else ""
                lines.append(f"- **{edu['titulo']}**{centro} ({edu['fecha']})")
            lines.append("")
        elif seccion == "idiomas":
            lines += ["## Idiomas", "", idiomas_linea(cv), ""]
        elif seccion == "certificaciones" and cv.get("certificaciones"):
            lines += ["## Certificaciones", "", certs_linea(cv), ""]
    return "\n".join(lines).rstrip() + "\n"


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), "Calibri")


def add_heading_docx(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    set_run_font(run, 11, bold=True)


def to_docx(cv: dict, dest: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(cv["nombre"])
    set_run_font(run, 18, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(cv["headline"])
    set_run_font(run, 12, bold=True)

    c = cv["contacto"]
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(f"{c['telefono']} · {c['email']} · {c['ciudad']}, {c['pais']}")
    set_run_font(run, 9.5)
    urls = contact_links(c)
    for i, url in enumerate(urls):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8) if i == len(urls) - 1 else Pt(1)
        run = p.add_run(url)
        set_run_font(run, 9)
    if not urls:
        p.paragraph_format.space_after = Pt(8)

    for seccion in cv.get("secciones") or []:
        if seccion == "perfil":
            add_heading_docx(doc, "Perfil profesional")
            p = doc.add_paragraph()
            run = p.add_run(cv.get("perfil") or "")
            set_run_font(run, 10.5)
        elif seccion == "competencias":
            add_heading_docx(doc, "Competencias")
            p = doc.add_paragraph()
            run = p.add_run(" · ".join(cv.get("competencias") or []))
            set_run_font(run, 10.5)
        elif seccion == "experiencia":
            add_heading_docx(doc, "Experiencia")
            for rol in cv.get("experiencia") or []:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(f"{rol['titulo']} — {rol['empresa']} ({rol['fechas']})")
                set_run_font(run, 10.5, bold=True)
                for b in rol.get("bullets") or []:
                    item = doc.add_paragraph(style="List Bullet")
                    run = item.add_run(bullet_text(b))
                    set_run_font(run, 10.5)
        elif seccion == "proyectos" and cv.get("proyectos"):
            add_heading_docx(doc, "Proyectos")
            for proj in cv["proyectos"]:
                p = doc.add_paragraph()
                url = f" — {proj['url']}" if proj.get("url") else ""
                run = p.add_run(f"{proj['nombre']}{url}")
                set_run_font(run, 10.5, bold=True)
                if proj.get("stack"):
                    p = doc.add_paragraph()
                    run = p.add_run(", ".join(proj["stack"]))
                    set_run_font(run, 9.5)
                for b in proj.get("bullets") or []:
                    item = doc.add_paragraph(style="List Bullet")
                    run = item.add_run(bullet_text(b))
                    set_run_font(run, 10.5)
        elif seccion == "formacion":
            add_heading_docx(doc, "Formación")
            for edu in cv.get("formacion") or []:
                centro = f" — {edu['centro']}" if edu.get("centro") else ""
                item = doc.add_paragraph(style="List Bullet")
                run = item.add_run(f"{edu['titulo']}{centro} ({edu['fecha']})")
                set_run_font(run, 10.5)
        elif seccion == "idiomas":
            add_heading_docx(doc, "Idiomas")
            p = doc.add_paragraph()
            run = p.add_run(idiomas_linea(cv))
            set_run_font(run, 10.5)
        elif seccion == "certificaciones" and cv.get("certificaciones"):
            add_heading_docx(doc, "Certificaciones")
            p = doc.add_paragraph()
            run = p.add_run(certs_linea(cv))
            set_run_font(run, 10.5)

    for p in doc.paragraphs:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def to_pdf(cv: dict, dest: Path) -> None:
    env = Environment(loader=FileSystemLoader(str(PLANTILLAS)), autoescape=True)
    template = env.get_template("cv.html.j2")
    css_path = PLANTILLAS / "cv.css"
    html = template.render(
        cv=cv,
        css_href=css_path.as_uri(),
        idiomas_linea=idiomas_linea(cv),
        certs_linea=certs_linea(cv),
        contact_links=contact_links(cv.get("contacto") or {}),
    )
    dest.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html, base_url=str(PLANTILLAS)).write_pdf(
        dest, stylesheets=[CSS(filename=str(css_path))]
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Renderiza cv.yaml a md, PDF y DOCX")
    parser.add_argument("--cv", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--basename", default="curriculum")
    args = parser.parse_args()

    cv = apply_perfil(load_yaml(args.cv))
    out = args.out_dir
    out.mkdir(parents=True, exist_ok=True)
    base = args.basename

    md_path = out / f"{base}.md"
    pdf_path = out / f"{base}.pdf"
    docx_path = out / f"{base}.docx"

    md_path.write_text(to_markdown(cv), encoding="utf-8")
    to_pdf(cv, pdf_path)
    to_docx(cv, docx_path)

    size = pdf_path.stat().st_size
    print(f"escrito {md_path}")
    print(f"escrito {pdf_path} ({size} bytes)")
    print(f"escrito {docx_path}")
    if size > 1_000_000:
        print("WARN  PDF supera 1 MB")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
